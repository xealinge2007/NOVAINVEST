"""Extractor de fundamentales para EEFF consolidados publicados en XLSX (no
PDF) por el propio emisor -- verificado contra un único caso real:
`ETB/2019-T3_Estados-Financieros-Consolidados-NIIF.xlsx`, que `ingesta_simev.py`
antes ignoraba por completo (el contrato de nombre y el escaneo de carpeta
solo cubrían `*.pdf`). Formato observado: tres hojas ('PYG - Income
Statement', 'Bce Gral - Bal Sh', 'Fluo Efectivo - Cash Flow'), etiquetas de
fila bilingües "Español / English" en la columna A, encabezados de columna
con el período (fecha puntual para el balance, "N ene - M mes AAAA" +
"NT AAAA" para resultados y flujo de caja).

No es un extractor XLSX universal -- si otro emisor publica un formato
distinto, esto fallará limpio (motivos explícitos, nunca un valor
inventado) y hay que ampliarlo o escribir uno nuevo, igual que
`extractor_generico.py` no cubre todo layout de PDF.

Contrato de salida idéntico a `extractor_generico.extraer()` para que
`extraer_fundamentales.py` no necesite una rama de código nueva más allá
de elegir qué extractor llamar:
    {"campos": {nombre_campo: {"valor": float|None, "pagina": None}},
     "unidad": str, "cuadra_balance": bool|None, "motivos": list[str]}
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

MESES_TRIMESTRE = {"T1": "Mar", "T2": "Jun", "T3": "Sep", "T4": "Dic"}

# (nombre_campo, fragmento de etiqueta a buscar en la columna A, mayus/minus
# insensible). Se usa el primer fragmento que aparezca en cada hoja -- ver
# `_buscar_fila`.
# Fragmentos deliberadamente largos/especificos -- verificado a mano que cada
# uno aparece UNA sola vez en la hoja real de ETB. Un fragmento corto como
# "Total de activos" matchea primero "Total de activos no corrientes" (fila
# de subtotal, valor mas chico) antes de llegar al total real -- ya paso una
# vez en pruebas contra este mismo archivo, balance no cuadraba por eso.
ETIQUETAS_BALANCE = [
    ("activos_totales", "Total de activos / Total assets"),
    ("pasivos_totales", "Total pasivos / Total Liabilities"),
    ("patrimonio", "Patrimonio total / Total Equity"),
]
ETIQUETAS_DEUDA = [
    "Obligaciones financieras corrientes",
    "Obligaciones financieras no corrientes",
]
ETIQUETAS_RESULTADOS = [
    ("ingresos", "Ingresos de Operaci"),
    ("utilidad_operacional", "Utilidad (Perdida) Recurrente / Recurring operating"),
    ("utilidad_neta", "Neta / Net income"),
]
ETIQUETA_FLUJO_OPERATIVO = "Flujos de efectivo netos procedentes de (utilizados en) actividades de operaci"


def _norm(s) -> str:
    return (s or "").replace("\n", " ").strip()


def _buscar_fila(ws, fragmento: str) -> list | None:
    """Primera fila cuya columna A contiene `fragmento` (insensible a
    mayúsculas/tildes exactas -- se compara la forma cruda, basta con que el
    fragmento venga sin acentos donde la fuente los rompe)."""
    frag = fragmento.lower()
    for row in ws.iter_rows(values_only=True):
        etiqueta = _norm(row[0])
        if frag in etiqueta.lower():
            return list(row)
    return None


def _columna_snapshot(ws, anio: int, periodo: str) -> int | None:
    """Índice (0-based) de la columna cuyo encabezado es la fecha puntual
    "A <mes> <año>" -- para la hoja de balance, donde cada columna es una
    foto a una fecha, nunca un acumulado."""
    mes = "Dic" if periodo == "ANUAL" else MESES_TRIMESTRE.get(periodo)
    if mes is None:
        return None
    patron = re.compile(rf"A\s+{mes}\s*{anio}", re.IGNORECASE)
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        for i, celda in enumerate(row):
            if i == 0:
                continue
            if patron.search(_norm(celda)):
                return i
    return None


def _columna_periodo(ws, anio: int, periodo: str) -> int | None:
    """Índice (0-based) de la columna del trimestre SUELTO ("NT AAAA") para
    resultados/flujo de caja, o de la columna acumulada anual ("1 ene - 31
    Dic AAAA") cuando `periodo == 'ANUAL'` -- ahí el acumulado del año
    completo YA ES el dato anual, no hace falta restar trimestres como en
    los PDF (`analizador_fundamental.py` resuelve eso para el PDF; este
    archivo trae el trimestre suelto de regalo)."""
    if periodo == "ANUAL":
        patron = re.compile(rf"1\s*ene.*31\s*Dic\s*{anio}", re.IGNORECASE)
    else:
        n = periodo[1]  # "T3" -> "3"
        patron = re.compile(rf"^{n}T\s*{anio}", re.IGNORECASE)
    for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
        for i, celda in enumerate(row):
            if i == 0:
                continue
            if patron.search(_norm(celda)):
                return i
    return None


def _valor(fila: list | None, col: int | None) -> float | None:
    if fila is None or col is None or col >= len(fila):
        return None
    v = fila[col]
    if v is None or not isinstance(v, (int, float)):
        return None
    return float(v)


# La unidad canonica del sistema es "miles_de_millones" (ver
# extractor_generico._detectar_factor_unidad) -- todo lo que entra a
# `fundamentales_reportados` esta en esa escala, y el guardia
# `_activos_fuera_de_rango` de extraer_fundamentales.py compara contra el
# historial del emisor asumiendola. Si esta hoja no se convierte al mismo
# factor, un valor correcto se ve "1000x mas grande" que el historial y se
# descarta como sospechoso (paso real con ETB: activos_totales en miles de
# pesos sin convertir disparo FUERA_DE_RANGO).
_FACTOR_A_MILES_DE_MILLONES = {
    "miles": 1e-6,       # miles de pesos -> miles de millones
    "millones": 1e-3,    # millones de pesos -> miles de millones
    "miles_de_millones": 1.0,
}


def extraer(ruta_xlsx: Path, anio: int, periodo: str) -> dict:
    motivos: list[str] = []
    campos: dict[str, dict] = {}
    cuadra_balance: bool | None = None
    unidad_detectada = None

    try:
        wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    except Exception as e:
        return {
            "campos": {},
            "unidad": None,
            "cuadra_balance": None,
            "motivos": [f"no se pudo abrir el XLSX: {type(e).__name__}: {e}"],
        }

    hoja_balance = next((wb[n] for n in wb.sheetnames if "bal" in n.lower() or "bce" in n.lower()), None)
    hoja_resultados = next((wb[n] for n in wb.sheetnames if "income" in n.lower() or "pyg" in n.lower()), None)
    hoja_flujo = next((wb[n] for n in wb.sheetnames if "cash" in n.lower() or "flujo" in n.lower()), None)

    if hoja_balance is not None:
        for fila_encabezado in hoja_balance.iter_rows(min_row=1, max_row=4, values_only=True):
            for celda in fila_encabezado:
                texto = str(celda).lower() if celda else ""
                if "miles de millones" in texto:
                    unidad_detectada = "miles_de_millones"
                elif "miles" in texto:
                    unidad_detectada = unidad_detectada or "miles"
                elif "millones" in texto:
                    unidad_detectada = unidad_detectada or "millones"

        col_balance = _columna_snapshot(hoja_balance, anio, periodo)
        if col_balance is None:
            motivos.append(f"hoja de balance: no se encontro columna para {periodo} {anio}")
        else:
            for campo, etiqueta in ETIQUETAS_BALANCE:
                fila = _buscar_fila(hoja_balance, etiqueta)
                valor = _valor(fila, col_balance)
                if valor is not None:
                    campos[campo] = {"valor": valor, "pagina": None}

            deuda_total = 0.0
            deuda_encontrada = False
            for etiqueta in ETIQUETAS_DEUDA:
                fila = _buscar_fila(hoja_balance, etiqueta)
                valor = _valor(fila, col_balance)
                if valor is not None:
                    deuda_total += valor
                    deuda_encontrada = True
            if deuda_encontrada:
                campos["deuda_financiera"] = {"valor": deuda_total, "pagina": None}
    else:
        motivos.append("no se encontro hoja de balance en el XLSX")

    if "activos_totales" in campos and "pasivos_totales" in campos and "patrimonio" in campos:
        suma = campos["pasivos_totales"]["valor"] + campos["patrimonio"]["valor"]
        cuadra_balance = abs(suma - campos["activos_totales"]["valor"]) <= 0.01 * campos["activos_totales"]["valor"]

    if hoja_resultados is not None:
        col_resultados = _columna_periodo(hoja_resultados, anio, periodo)
        if col_resultados is None:
            motivos.append(f"hoja de resultados: no se encontro columna para {periodo} {anio}")
        else:
            for campo, etiqueta in ETIQUETAS_RESULTADOS:
                fila = _buscar_fila(hoja_resultados, etiqueta)
                valor = _valor(fila, col_resultados)
                if valor is not None:
                    campos[campo] = {"valor": valor, "pagina": None}

    if hoja_flujo is not None:
        col_flujo = _columna_periodo(hoja_flujo, anio, periodo)
        if col_flujo is not None:
            fila = _buscar_fila(hoja_flujo, ETIQUETA_FLUJO_OPERATIVO)
            valor = _valor(fila, col_flujo)
            if valor is not None:
                campos["flujo_caja_operativo"] = {"valor": valor, "pagina": None}

    if not campos and not motivos:
        motivos.append("ninguna etiqueta de fila conocida coincidio -- formato XLSX distinto al verificado (ETB)")

    unidad_final = None
    if campos:
        if unidad_detectada is None:
            motivos.append("unidad no declarada en el encabezado de la hoja de balance -- valores sin convertir, NO publicar tal cual")
        else:
            factor = _FACTOR_A_MILES_DE_MILLONES[unidad_detectada]
            for c in campos.values():
                c["valor"] = c["valor"] * factor
            unidad_final = "miles_de_millones"

    return {
        "campos": campos if unidad_final else {},
        "unidad": unidad_final,
        "cuadra_balance": cuadra_balance,
        "motivos": motivos,
    }
